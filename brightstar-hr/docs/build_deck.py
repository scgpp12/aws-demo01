# -*- coding: utf-8 -*-
"""BrightStar 社内システム ご紹介デッキ（python-pptx → pptx）。

pptx 生成後、sons02 の LibreOffice で pdf 変換する：
  soffice --headless --convert-to pdf プレゼン資料.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY = RGBColor(0x0E, 0x22, 0x38)
NAVY2 = RGBColor(0x16, 0x31, 0x4E)
TEAL = RGBColor(0x13, 0x90, 0xA6)
TEALX = RGBColor(0x16, 0xB5, 0xCF)
GOLD = RGBColor(0xE0, 0xA4, 0x3B)
INK = RGBColor(0x1B, 0x27, 0x33)
MUTE = RGBColor(0x5E, 0x6B, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MIST = RGBColor(0xF4, 0xF7, 0xFB)

FONT = "Noto Sans JP"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height


def _rect(slide, x, y, w, h, color, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
    sp.shadow.inherit = False
    return sp


def _text(slide, x, y, w, h, runs, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, space=6):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [(runs, size, color, bold)]
    for i, r in enumerate(runs):
        txt, sz, col, bd = (r + (size, color, bold))[:4] if isinstance(r, tuple) else (r, size, color, bold)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space)
        run = p.add_run(); run.text = txt
        f = run.font; f.size = Pt(sz); f.bold = bd; f.color.rgb = col; f.name = FONT
    return tb


def content_slide(title, kicker, bullets, accent=TEAL):
    s = prs.slides.add_slide(BLANK)
    _rect(s, 0, 0, W, Inches(1.25), NAVY)
    _rect(s, 0, Inches(1.25), W, Inches(0.06), accent)
    _text(s, Inches(0.6), Inches(0.18), W - Inches(1.2), Inches(0.4),
          kicker, size=12, color=TEALX, bold=True)
    _text(s, Inches(0.6), Inches(0.5), W - Inches(1.2), Inches(0.7),
          title, size=26, color=WHITE, bold=True)
    body = []
    for b in bullets:
        if isinstance(b, tuple):
            body.append(b)
        else:
            body.append(("・" + b, 18, INK, False))
    _text(s, Inches(0.8), Inches(1.6), W - Inches(1.6), H - Inches(2.0),
          body, size=18, color=INK, space=10)
    _text(s, Inches(0.6), H - Inches(0.5), W - Inches(1.2), Inches(0.4),
          "BrightStar 社内システム", size=10, color=MUTE)
    return s


# 1. 表紙
s = prs.slides.add_slide(BLANK)
_rect(s, 0, 0, W, H, NAVY)
_rect(s, 0, Inches(3.0), W, Inches(0.08), GOLD)
_text(s, Inches(0.9), Inches(1.2), W - Inches(1.8), Inches(0.5),
      "CASE STUDY ／ サーバーレス × LINE × CDK", size=14, color=GOLD, bold=True)
_text(s, Inches(0.9), Inches(1.7), W - Inches(1.8), Inches(1.4),
      "BrightStar 社内システム", size=46, color=WHITE, bold=True)
_text(s, Inches(0.9), Inches(3.3), W - Inches(1.8), Inches(2.0), [
    ("勤怠（作業時間記録簿）と通勤費（経費）の提出・督促・集計を、", 22, WHITE, False),
    ("LINE のチャットだけで完結。", 22, TEALX, True),
    ("社員名簿を主データに未提出者を可視化し、自動リマインド・一括DLまで。", 18, RGBColor(0xCF, 0xDC, 0xE8), False),
    ("AWS サーバーレス ＋ AWS CDK(TypeScript) でフルスクラッチ開発。", 18, RGBColor(0xCF, 0xDC, 0xE8), False),
], space=10)
_text(s, Inches(0.9), H - Inches(0.9), W - Inches(1.8), Inches(0.5),
      "受託開発ケーススタディ ｜ 2026", size=12, color=RGBColor(0x9F, 0xB2, 0xC4))

# 2. 課題
content_slide("課題", "PROBLEM", [
    "毎月の勤怠表・通勤費精算が、メール添付・紙・表計算でバラバラ。",
    "「誰が出していないか」が見えず、人事が一人ずつ催促。",
    "専用アプリ導入・ID 配布は重い。社員は普段使う LINE で済ませたい。",
    "提出物の年月間違い・様式違いが後で発覚し、差し戻しが発生。",
], accent=GOLD)

# 3. ソリューション
content_slide("ソリューション", "SOLUTION", [
    "社員：記入した Excel を LINE に送るだけで提出（様式取得・再提出・履歴も LINE）。",
    "人事：LINE で全社の提出状況を即確認、未提出者へワンタップ督促、当月分を一括DL。",
    "名簿主導：社員名簿（DB）が「提出対象」を決定。未登録者も ※Line未登録 で可視化。",
    "提出時に年月を自動チェックし、月違いファイルはアップロードを拒否。",
])

# 4. 主要機能
content_slide("主要機能", "FEATURES", [
    ("■ 社員", 19, TEAL, True),
    "提出（再提出可）／テンプレ取得／個人履歴（ダウンロード付き）",
    "氏名だけで登録（部署は名簿から自動取得）",
    ("■ 人事", 19, TEAL, True),
    "未提出確認／一覧（未登録者も可視化）／催促／一括DL",
    "社員名簿の追加・変更・削除（変更は人事全員へ自動通知）",
])

# 5. 自動リマインド
content_slide("自動・手動リマインド", "REMINDER", [
    "自動：毎月 25・28 日 09:00（JST）に未提出者だけへ自動配信。",
    "手動：人事が「催促」で即時督促。",
    "提出済みは自動的に対象外（提出＝記録更新で差分から除外）。",
    "LINE 未登録の未提出者は一覧に ※Line未登録 を表示し、人事が個別連絡。",
])

# 6. 内容チェック（年月一致）
content_slide("内容チェック（年月の自動照合）", "VALIDATION", [
    "経費：セル A1「YYYY年M月分経費」の年月を読み取り。",
    "勤怠：セル B5 の日付（YYYY/M/D）の年月を読み取り。",
    "提出月と一致しない場合はアップロードを拒否し、社員へその場で通知。",
    "標準ライブラリのみで xlsx を解析（依存ゼロ／日付シリアルも対応）。",
], accent=GOLD)

# 7. アーキテクチャ
content_slide("フルサーバーレス構成（CDK）", "ARCHITECTURE", [
    "LINE 公式アカウント → Lambda Function URL → webhook(Python3.12/arm64)",
    "DynamoDB ×3（roster 名簿 / employees 連携 / submissions 提出）＋ S3（提出物・様式・zip）",
    "EventBridge（cron）→ reminder Lambda → 未提出者へ LINE push",
    "LINE は署名検証のみで接続（中継・暗号・可信IP 不要）。固定サーバなし・従量課金。",
])

# 8. セキュリティ・データ保護
content_slide("セキュリティ・データ保護", "SECURITY", [
    "LINE 資格情報は SSM SecureString（git・環境変数平文に置かない）。",
    "S3 は Block Public Access ＋ 暗号化 ＋ バージョン管理（再提出の旧版も保持）。",
    "ダウンロードは HMAC 署名付き短縮リンクで越権防止。",
    "保管：2 か月は即時DL、以降 Deep Archive（最安）、1 年で自動削除。",
])

# 9. 効果
content_slide("導入効果", "IMPACT", [
    "提出・督促・集計の工数を大幅削減、未提出の取りこぼしゼロへ。",
    "専用アプリ不要（LINE）。固定費ほぼゼロの従量運用。",
    "年月・様式の自動チェックで差し戻しを削減。",
    "CDK 管理で変更・横展開が容易（事業所追加・項目追加など）。",
])

# 10. ロードマップ
content_slide("拡張ロードマップ（任意）", "ROADMAP", [
    "Excel 中身の自動集計（工数・金額）。",
    "配信時刻の人事セルフ設定／複数事業所・多言語強化。",
    "他チャネル（Slack / Teams 等）への展開。",
])

# 11. クロージング
s = prs.slides.add_slide(BLANK)
_rect(s, 0, 0, W, H, NAVY)
_text(s, Inches(0.9), Inches(2.6), W - Inches(1.8), Inches(1.2),
      "同様の社内 DX を、御社にも。", size=36, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
_text(s, Inches(0.9), Inches(3.9), W - Inches(1.8), Inches(0.8),
      "メッセージング連携・サーバーレス基盤・IaC（CDK）の受託開発を承ります。",
      size=18, color=RGBColor(0xCF, 0xDC, 0xE8), align=PP_ALIGN.CENTER)
_rect(s, Inches(5.4), Inches(5.0), Inches(2.5), Inches(0.06), GOLD)

OUT = "プレゼン資料.pptx"
prs.save(OUT)
print("saved", OUT, "slides=", len(prs.slides._sldIdLst))
